#!/usr/bin/env python
"""Generate FraudShield Technical Report as PDF and DOCX from one content source.

Run from the repo root:  .venv/Scripts/python tools_submission/build_report.py
Writes submission/FraudShield_Technical_Report.pdf and .docx
"""
import os
from docx import Document
from docx.shared import Pt, RGBColor
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.lib import colors
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable,
)

OUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "submission"))
os.makedirs(OUT_DIR, exist_ok=True)

BACKEND_URL = "https://fraudshield-production-968e.up.railway.app"
FRONTEND_URL = "https://frontend-ruddy-chi-42.vercel.app"
REPO_URL = "https://github.com/Shayan121363/FraudShield"

TITLE = "FraudShield — Real-Time Transaction Fraud Detection"
SUBTITLE = "Technical Report · Alibaba Cloud AI Hackathon (Qualification Round)"

# ("h", text) heading | ("p", text) paragraph | ("b", text) bullet
# ("table", [[header...], [row...], ...])
SECTIONS = [
    ("p", "FraudShield is an end-to-end fraud detection system that scores every card "
          "transaction with a supervised model (XGBoost) and an unsupervised anomaly "
          "detector (a PyTorch autoencoder), fuses them into a single risk score, explains "
          "each decision with SHAP, and streams results to a live analyst console over "
          "WebSocket. A bilingual Confidence Flow layer turns those signals into plain-language "
          "safety guidance in English and Urdu. The system is deployed and publicly reachable."),
    ("table", [
        ["Live endpoints", ""],
        ["Analyst console (frontend)", FRONTEND_URL],
        ["API backend (REST + WebSocket)", BACKEND_URL],
        ["API health probe", BACKEND_URL + "/health"],
        ["Source repository", REPO_URL],
    ]),

    ("h", "1. The Problem"),
    ("p", "Card fraud is hard to model for reasons that go beyond 'train a classifier'. "
          "FraudShield is engineered around four of them."),
    ("b", "Severe class imbalance — fraud is a fraction of a percent of traffic, so accuracy "
          "is meaningless; the system reports PR-AUC and precision/recall at a tuned threshold."),
    ("b", "Concept drift — fraud rings change tactics quickly; a supervised-only model can only "
          "find fraud that resembles already-labelled fraud, so a second unsupervised model "
          "that never sees a fraud label catches novel patterns."),
    ("b", "Black-box scores are not actionable — an analyst cannot block a payment because "
          "'the model said 0.91'. Every score is attributed to its driving features and rendered "
          "as a written justification."),
    ("b", "Financial inclusion — a blocked payment must also be explained to the customer in "
          "plain language, especially first-time digital banking users, which motivated the "
          "bilingual English/Urdu Confidence Flow."),

    ("h", "2. System Architecture"),
    ("p", "A transaction (POST /predict, or sampled from the simulated stream) is scaled with a "
          "fitted StandardScaler and scored by two models in parallel: XGBoost (supervised "
          "probability) and the autoencoder (reconstruction-error anomaly score). The two are "
          "combined by a weighted ensemble (0.7 supervised / 0.3 anomaly), then SHAP TreeExplainer "
          "attributes the result to the top features. The enriched object — three scores, a risk "
          "band, top factors, an analyst explanation and bilingual user guidance — is returned over "
          "REST, pushed to the dashboard over WebSocket, and persisted to the audit store, so what "
          "an analyst sees and what the audit trail records cannot drift apart."),
    ("b", "Backend: FastAPI, Pydantic v2, SQLAlchemy 2, Uvicorn, WebSockets."),
    ("b", "Data: PostgreSQL in production, SQLite locally — selected by a single DATABASE_URL "
          "variable with no divergent code path."),
    ("b", "Frontend: React 19 + Vite 8, Recharts, React Router (hash routing)."),

    ("h", "3. Machine-Learning Approach"),
    ("table", [
        ["Component", "Technique", "Rationale"],
        ["Class imbalance", "SMOTE oversampling (training split only)", "Prevents collapse to the majority class; test set is never resampled"],
        ["Supervised model", "XGBoost", "Strong on tabular financial data, fast, supports exact tree-path SHAP"],
        ["Unsupervised model", "PyTorch autoencoder (7-16-8-3-8-16-7)", "Trained only on legitimate transactions; scores by reconstruction error"],
        ["Fusion", "Weighted ensemble (0.7 / 0.3)", "The two models fail differently, so the pair is more robust"],
        ["Explainability", "SHAP TreeExplainer", "Additive, per-transaction, theoretically consistent attributions"],
        ["Analyst reporting", "Rule-based generation from SHAP values", "Deterministic and auditable"],
        ["User guidance", "Rule-based bilingual templates (EN/UR)", "Plain-language safety checklists for underserved users"],
        ["Evaluation", "PR-AUC, precision/recall, F1-optimal threshold", "Accuracy is meaningless under imbalance"],
    ]),

    ("h", "4. Results"),
    ("table", [
        ["Metric", "Value"],
        ["XGBoost ROC-AUC", "0.9999"],
        ["XGBoost PR-AUC", "0.9995"],
        ["Autoencoder ROC-AUC (reconstruction error)", "0.9929"],
        ["Ensemble PR-AUC", "0.9984"],
        ["F1-optimal decision threshold (ensemble)", "0.5719"],
        ["Anomaly calibration threshold (95th pct legit)", "0.10117"],
        ["Untuned LogisticRegression PR-AUC, same corpus", "0.9996"],
    ]),
    ("p", "Honest interpretation: these numbers come from a 50,000-row synthetic corpus generated "
          "at a 10% fraud rate whose classes are more separated than real traffic. The final row is "
          "the key caveat — a seven-feature untuned logistic regression edges past the ensemble on "
          "this data, and three of the seven individual features are clear 0.8 PR-AUC on their own. "
          "Read these metrics as proof that the architecture and explainability layer work end to "
          "end, and NOT as evidence of real-world detection quality. Re-baselining on the real "
          "IEEE-CIS corpus (590,540 transactions, 20,663 frauds) with a temporal split and "
          "leak-free merchant-risk encoding is the next step, and the numbers will change."),

    ("h", "5. Explainability & Bilingual Confidence Flow"),
    ("p", "Every scoring response includes the top SHAP factors, an analyst-facing explanation, and "
          "a structured user_guidance block: a demographic label, plain-language title/reason and a "
          "recommended action in both English and Urdu, plus an interactive safety checklist "
          "(OTP/PIN sharing, recipient verification, and similar). This is the project's financial-"
          "inclusion angle — decisions are explained to end users, not just scored."),

    ("h", "6. API Surface"),
    ("table", [
        ["Endpoint", "Method", "Description"],
        ["/predict", "POST", "Score one transaction"],
        ["/predict/batch", "POST", "Score a list of transactions"],
        ["/history", "GET", "Last N scored transactions from the database, newest first"],
        ["/stats", "GET", "Session totals, flagged count, fraud rate, uptime"],
        ["/health", "GET", "Liveness, loaded metrics, database driver"],
        ["/ws/stream", "WebSocket", "Live simulated feed of pre-scored transactions"],
    ]),

    ("h", "7. Deployment"),
    ("p", "The backend runs as a Dockerized FastAPI/Uvicorn service on Railway (CPU-only PyTorch, "
          "libgomp1 for torch on the slim base image) with a managed PostgreSQL audit store. The "
          "React console is built with Vite and hosted on Vercel, configured at build time to call "
          "the Railway backend over HTTPS and WSS. Both are live and public (see URLs above)."),
    ("b", "Model artifacts are committed to the repo so a fresh clone runs out of the box."),
    ("b", "Environment is selected entirely by variables: DATABASE_URL on the backend, "
          "VITE_API_URL / VITE_WS_URL baked into the frontend build."),

    ("h", "8. Reproducibility & Verification"),
    ("p", "train.py is deterministic (pinned seeds, single-threaded XGBoost, anomaly threshold "
          "calibrated on legitimate training rows only), producing byte-identical artifacts across "
          "runs. scripts/verify_pipeline.py asserts the serving-path contract end to end (12 checks) "
          "and probes the linear baseline. The committed metrics.json is served verbatim by /health."),

    ("h", "9. Known Limitations & Roadmap"),
    ("b", "Headline metrics are synthetic-data artifacts; the ensemble's 0.3 anomaly weight is "
          "asserted rather than proven on this data."),
    ("b", "No endpoint authentication and CORS is open (fine for a demo, not for production)."),
    ("b", "SESSION_STATS is process-local; durable figures live in /history."),
    ("p", "Roadmap: (1) retrain on IEEE-CIS with a temporal, leak-free split; (2) Alibaba Cloud "
          "deployment with artifacts in OSS under versioned keys and managed PostgreSQL; (3) analyst "
          "review queue that also produces the next training labels; (4) auth and input hardening; "
          "(5) optional Qwen (Model Studio) narrative reports grounded on SHAP values."),
]


def _pdf():
    styles = getSampleStyleSheet()
    h1 = ParagraphStyle("H1", parent=styles["Heading1"], fontSize=20, leading=24,
                        textColor=colors.HexColor("#0B3D91"))
    sub = ParagraphStyle("Sub", parent=styles["Normal"], fontSize=11,
                        textColor=colors.HexColor("#555555"))
    h2 = ParagraphStyle("H2", parent=styles["Heading2"], fontSize=13, spaceBefore=10,
                        textColor=colors.HexColor("#0B3D91"))
    body = ParagraphStyle("Body", parent=styles["Normal"], fontSize=10, leading=14,
                          spaceAfter=6)
    bullet = ParagraphStyle("Bullet", parent=body, leftIndent=12, bulletIndent=2)

    story = [Paragraph(TITLE, h1), Paragraph(SUBTITLE, sub), Spacer(1, 6),
             HRFlowable(width="100%", color=colors.HexColor("#0B3D91")), Spacer(1, 10)]

    for kind, payload in SECTIONS:
        if kind == "h":
            story.append(Paragraph(payload, h2))
        elif kind == "p":
            story.append(Paragraph(payload, body))
        elif kind == "b":
            story.append(Paragraph("• " + payload, bullet))
        elif kind == "table":
            data = [[Paragraph(str(c), body) for c in row] for row in payload]
            tbl = Table(data, colWidths=[45 * mm, 55 * mm, 65 * mm][:len(payload[0])]
                        if len(payload[0]) == 3 else None)
            tbl.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0B3D91")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BBBBBB")),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("FONTSIZE", (0, 0), (-1, -1), 9),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F2F6FF")]),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
            ]))
            story.append(tbl)
            story.append(Spacer(1, 8))

    path = os.path.join(OUT_DIR, "FraudShield_Technical_Report.pdf")
    SimpleDocTemplate(path, pagesize=A4,
                      leftMargin=20 * mm, rightMargin=20 * mm,
                      topMargin=18 * mm, bottomMargin=18 * mm,
                      title=TITLE).build(story)
    return path


def _docx():
    doc = Document()
    doc.styles["Normal"].font.name = "Calibri"
    doc.styles["Normal"].font.size = Pt(11)

    t = doc.add_heading(TITLE, level=0)
    for r in t.runs:
        r.font.color.rgb = RGBColor(0x0B, 0x3D, 0x91)
    p = doc.add_paragraph(SUBTITLE)
    p.runs[0].font.color.rgb = RGBColor(0x66, 0x66, 0x66)
    doc.add_paragraph()

    for kind, payload in SECTIONS:
        if kind == "h":
            h = doc.add_heading(payload, level=1)
            for r in h.runs:
                r.font.color.rgb = RGBColor(0x0B, 0x3D, 0x91)
        elif kind == "p":
            doc.add_paragraph(payload)
        elif kind == "b":
            doc.add_paragraph(payload, style="List Bullet")
        elif kind == "table":
            rows = len(payload)
            cols = len(payload[0])
            table = doc.add_table(rows=rows, cols=cols)
            table.style = "Light Grid Accent 1"
            for i, row in enumerate(payload):
                for j, cell in enumerate(row):
                    table.cell(i, j).text = str(cell)
                    if i == 0:
                        for r in table.cell(i, j).paragraphs[0].runs:
                            r.bold = True
            doc.add_paragraph()

    path = os.path.join(OUT_DIR, "FraudShield_Technical_Report.docx")
    doc.save(path)
    return path


if __name__ == "__main__":
    print("PDF :", _pdf())
    print("DOCX:", _docx())
