#!/usr/bin/env python
"""Generate FraudShield Pitch Deck (.pptx) for the Alibaba Cloud AI Hackathon.

Run:  .venv/Scripts/python tools_submission/build_deck.py
Writes submission/FraudShield_Pitch_Deck.pptx
"""
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "submission"))
os.makedirs(OUT_DIR, exist_ok=True)

NAVY = RGBColor(0x0B, 0x3D, 0x91)
GREY = RGBColor(0x55, 0x55, 0x55)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
ACCENT = RGBColor(0xE8, 0x4F, 0x4F)  # risk red

BACKEND_URL = "https://fraudshield-production-968e.up.railway.app"
FRONTEND_URL = "https://frontend-ruddy-chi-42.vercel.app"
REPO_URL = "https://github.com/Shayan121363/FraudShield"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW, SH = prs.slide_width, prs.slide_height


def add_slide():
    return prs.slides.add_slide(BLANK)


def fill_bg(slide, color):
    shp = slide.shapes.add_shape(1, 0, 0, SW, SH)  # rectangle
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    slide.shapes._spTree.remove(shp._element)
    slide.shapes._spTree.insert(2, shp._element)
    return shp


def textbox(slide, l, t, w, h):
    tb = slide.shapes.add_textbox(l, t, w, h)
    tf = tb.text_frame
    tf.word_wrap = True
    return tf


def set_para(p, text, size, color, bold=False, align=PP_ALIGN.LEFT):
    p.text = text
    p.alignment = align
    run = p.runs[0]
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def bullets(slide, l, t, w, h, items, size=18, color=GREY, gap=6):
    tf = textbox(slide, l, t, w, h)
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        level_lead = "▸  "
        set_para(p, level_lead + it, size, color)
        p.space_after = Pt(gap)
    return tf


# ---- Slide 1: Title ----
s = add_slide()
fill_bg(s, NAVY)
tf = textbox(s, Inches(1), Inches(2.4), Inches(11.3), Inches(2))
set_para(tf.paragraphs[0], "FraudShield", 54, WHITE, bold=True)
p = tf.add_paragraph()
set_para(p, "Real-Time Transaction Fraud Detection", 26, RGBColor(0xC9, 0xD8, 0xFF))
p = tf.add_paragraph()
p.space_before = Pt(20)
set_para(p, "Alibaba Cloud AI Hackathon · Qualification Round", 16, WHITE)

# ---- Slide 2: Problem ----
s = add_slide()
tf = textbox(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
set_para(tf.paragraphs[0], "The Problem", 34, NAVY, bold=True)
bullets(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5), [
    "Severe class imbalance — fraud is a fraction of a percent; accuracy is misleading.",
    "Concept drift — supervised-only models miss novel attacks exactly when losses begin.",
    "Black-box scores aren't actionable — analysts and compliance need a justification.",
    "Financial inclusion — blocked payments must be explained to everyday customers.",
], size=22, color=GREY, gap=14)

# ---- Slide 3: Solution ----
s = add_slide()
tf = textbox(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
set_para(tf.paragraphs[0], "The Solution", 34, NAVY, bold=True)
bullets(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5), [
    "Ensemble scoring: XGBoost (supervised) + PyTorch autoencoder (anomaly), 0.7 / 0.3 fusion.",
    "Explain every decision with SHAP and a written analyst explanation.",
    "Bilingual Confidence Flow: plain-language safety guidance in English and Urdu.",
    "Live analyst console streamed over WebSocket, with a durable audit trail.",
], size=22, color=GREY, gap=14)

# ---- Slide 4: Architecture ----
s = add_slide()
tf = textbox(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
set_para(tf.paragraphs[0], "How It Works", 34, NAVY, bold=True)
steps = [
    ("Transaction", "POST /predict or live stream"),
    ("StandardScaler", "normalize 7 features"),
    ("XGBoost + Autoencoder", "probability + anomaly"),
    ("Ensemble + SHAP", "risk score + attribution"),
    ("Console · API · DB", "WS feed, REST, audit store"),
]
x = Inches(0.7)
w = Inches(2.25)
for i, (title, sub) in enumerate(steps):
    box = s.shapes.add_shape(1, x, Inches(2.6), w, Inches(1.7))
    box.fill.solid()
    box.fill.fore_color.rgb = NAVY if i % 2 == 0 else RGBColor(0x1E, 0x57, 0xB0)
    box.line.fill.background()
    tfx = box.text_frame
    tfx.word_wrap = True
    set_para(tfx.paragraphs[0], title, 15, WHITE, bold=True, align=PP_ALIGN.CENTER)
    p = tfx.add_paragraph()
    set_para(p, sub, 11, RGBColor(0xD6, 0xE0, 0xFF), align=PP_ALIGN.CENTER)
    x += w + Inches(0.18)
bullets(s, Inches(0.9), Inches(4.9), Inches(11.5), Inches(1.6), [
    "One scored object feeds REST, WebSocket and the database, so UI and audit never diverge.",
    "Stack: FastAPI · Pydantic v2 · SQLAlchemy 2 · React 19 · Vite · PostgreSQL/SQLite.",
], size=16, color=GREY, gap=6)

# ---- Slide 5: Results (with honesty) ----
s = add_slide()
tf = textbox(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
set_para(tf.paragraphs[0], "Results (and honest context)", 34, NAVY, bold=True)
rows = [
    ("Model", "Metric", "Value"),
    ("XGBoost", "ROC-AUC / PR-AUC", "0.9999 / 0.9995"),
    ("Autoencoder", "ROC-AUC", "0.9929"),
    ("Ensemble", "PR-AUC", "0.9984"),
    ("Baseline", "LogisticRegression PR-AUC", "0.9996"),
]
tbl_shape = s.shapes.add_table(len(rows), 3, Inches(0.9), Inches(1.7),
                               Inches(8.2), Inches(2.6))
tbl = tbl_shape.table
for i, row in enumerate(rows):
    for j, val in enumerate(row):
        cell = tbl.cell(i, j)
        cell.text = val
        para = cell.text_frame.paragraphs[0]
        para.runs[0].font.size = Pt(15)
        para.runs[0].font.bold = (i == 0)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY if i == 0 else WHITE
        para.runs[0].font.color.rgb = WHITE if i == 0 else GREY
bullets(s, Inches(0.9), Inches(4.7), Inches(11.5), Inches(2), [
    "Metrics are on a 50k synthetic corpus at a 10% fraud rate — saturated and 'too easy'.",
    "An untuned logistic regression matches the ensemble here: architecture is proven, detection quality is not yet.",
    "Next: retrain on IEEE-CIS (590k txns) with a temporal, leak-free split.",
], size=15, color=ACCENT, gap=6)

# ---- Slide 6: Social impact / inclusion ----
s = add_slide()
tf = textbox(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
set_para(tf.paragraphs[0], "Inclusion & Explainability", 34, NAVY, bold=True)
bullets(s, Inches(0.9), Inches(1.7), Inches(11.5), Inches(5), [
    "Every score ships with SHAP-driven top factors and a written explanation.",
    "Confidence Flow translates signals into an interactive safety checklist.",
    "Fully bilingual English / Urdu with right-to-left layout for first-time digital users.",
    "Turns a model verdict into something a customer can actually act on.",
], size=22, color=GREY, gap=14)

# ---- Slide 7: Live demo ----
s = add_slide()
fill_bg(s, RGBColor(0xF2, 0xF6, 0xFF))
tf = textbox(s, Inches(0.8), Inches(0.5), Inches(11.7), Inches(1))
set_para(tf.paragraphs[0], "Live & Deployed", 34, NAVY, bold=True)
tf2 = textbox(s, Inches(0.9), Inches(1.8), Inches(11.5), Inches(4))
set_para(tf2.paragraphs[0], "Analyst console (Vercel)", 20, NAVY, bold=True)
p = tf2.add_paragraph(); set_para(p, FRONTEND_URL, 20, GREY)
p = tf2.add_paragraph(); p.space_before = Pt(14); set_para(p, "API backend (Railway)", 20, NAVY, bold=True)
p = tf2.add_paragraph(); set_para(p, BACKEND_URL, 20, GREY)
p = tf2.add_paragraph(); p.space_before = Pt(14); set_para(p, "Health probe", 20, NAVY, bold=True)
p = tf2.add_paragraph(); set_para(p, BACKEND_URL + "/health", 20, GREY)
p = tf2.add_paragraph(); p.space_before = Pt(14); set_para(p, "Repository", 20, NAVY, bold=True)
p = tf2.add_paragraph(); set_para(p, REPO_URL, 20, GREY)

# ---- Slide 8: Roadmap / thanks ----
s = add_slide()
fill_bg(s, NAVY)
tf = textbox(s, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.5))
set_para(tf.paragraphs[0], "Roadmap", 34, WHITE, bold=True)
for i, item in enumerate([
    "Retrain on real IEEE-CIS data (temporal split, leak-free features)",
    "Alibaba Cloud deployment — containers, OSS-versioned artifacts, managed PostgreSQL",
    "Analyst review queue that feeds the next training labels",
    "Endpoint auth + input hardening; optional Qwen narrative reports",
]):
    p = tf.add_paragraph(); p.space_before = Pt(10)
    set_para(p, "▸  " + item, 20, RGBColor(0xD6, 0xE0, 0xFF))

path = os.path.join(OUT_DIR, "FraudShield_Pitch_Deck.pptx")
prs.save(path)
print("PPTX:", path)
